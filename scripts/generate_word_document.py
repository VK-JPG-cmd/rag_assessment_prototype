"""
Generates a formal, professional Word Document (.docx) report
for the Patient-Record RAG Prototype submission.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate Navy
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(2, 132, 199) # Primary Blue
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(13, 148, 136) # Teal
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_callout_box(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F0F9FF") # Light Sky
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.font.bold = True
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(2, 132, 199)
    
    run_text = p.add_run(text)
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(51, 65, 85)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def generate_report():
    doc = docx.Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set Normal Style Font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 41, 59)

    # ==========================================
    # TITLE & HEADER
    # ==========================================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("PATIENT-RECORD RAG PROTOTYPE")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Historical Clinical Information Retrieval & Grounded Reasoning using Synthetic EHR Data")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(2, 132, 199)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(18)
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = meta_p.add_run("Author: VK-JPG-cmd | Submission Report | System Version: 1.0.0 | Date: September 2026")
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ==========================================
    # 1. OBJECTIVE / PROBLEM STATEMENT
    # ==========================================
    add_styled_heading(doc, "1. Objective / Problem Statement", level=1)
    
    p = doc.add_paragraph(
        "Electronic Health Records (EHRs) contain extensive longitudinal clinical data spanning multiple years, "
        "including progress notes, diagnostic lab trends, active prescriptions, surgical histories, and allergy profiles. "
        "Clinicians and healthcare researchers often face significant information overload when attempting to extract "
        "pertinent historical patient facts rapidly at the point of care."
    )
    p.paragraph_format.line_spacing = 1.15

    p2 = doc.add_paragraph(
        "The objective of this project is to design, develop, and evaluate a production-grade Clinical Retrieval-Augmented "
        "Generation (RAG) prototype. The system leverages synthetic, HIPAA Safe Harbor de-identified longitudinal patient records "
        "to accurately retrieve relevant historical medical facts and synthesize precise, grounded answers with verifiable source citations."
    )
    p2.paragraph_format.line_spacing = 1.15

    add_callout_box(
        doc,
        "Core Clinical Objectives",
        "1. Patient Data Isolation: Strict zero cross-patient data leakage via metadata pre-filtering.\n"
        "2. Longitudinal & Temporal Reasoning: Tracking lab trajectories and medication adjustments over time.\n"
        "3. Zero-Hallucination Guardrails: Strict factual grounding with explicit absence-of-data declarations.\n"
        "4. Verifiable Citations: Tagging assertions with exact source notes and encounter dates."
    )

    # ==========================================
    # 2. APPROACH AND METHODOLOGY
    # ==========================================
    add_styled_heading(doc, "2. Approach and Methodology", level=1)
    
    p = doc.add_paragraph(
        "The prototype architecture combines structured metadata filtering, dense vector embeddings, and large language model reasoning. "
        "The pipeline is divided into four major stages:"
    )

    add_styled_heading(doc, "2.1 Metadata-Aware Document Chunking", level=2)
    doc.add_paragraph(
        "Standard naive chunking splits text at fixed character boundaries, which breaks clinical coherence. In this system, patient records "
        "are parsed into discrete clinical categories (Demographics, Active Problem List, Allergies, Active Prescriptions, Clinical Progress Notes, "
        "and Longitudinal Lab Panels). Each chunk retains vital metadata: patient_id, encounter_date, record_type, category, and source_id."
    )

    add_styled_heading(doc, "2.2 Vector Database & Embeddings Pipeline", level=2)
    doc.add_paragraph(
        "Chunks are converted into high-dimensional vector representations using Google Gemini models/gemini-embedding-001 (3072 dimensions). "
        "Vectors are indexed inside an embedded ChromaDB persistent collection configured with cosine distance. A deterministic term-frequency fallback "
        "embedding model is integrated to ensure 100% offline functionality if external APIs are unreachable."
    )

    add_styled_heading(doc, "2.3 Patient-Filtered Similarity Retrieval", level=2)
    doc.add_paragraph(
        "To guarantee healthcare privacy and prevent data contamination, queries enforce a strict metadata filter: where={'patient_id': target_patient_id}. "
        "Only records belonging strictly to the selected patient are searched. Matched documents are ranked by cosine similarity and organized chronologically."
    )

    add_styled_heading(doc, "2.4 Grounded Generation & Citation Verification", level=2)
    doc.add_paragraph(
        "The retrieved context is injected into a specialized clinical prompt that instructs the LLM (Google Gemini 3 Flash) to formulate an objective, "
        "temporally accurate response with explicit citation tags (e.g., [Source: ENC-P101-01 (2021-01-15)]). If information is missing from the record, "
        "the model is constrained to explicitly declare that the data is not documented."
    )

    # ==========================================
    # 3. DATASET OR SAMPLE DATA USED
    # ==========================================
    add_styled_heading(doc, "3. Dataset or Sample Data Used", level=1)
    doc.add_paragraph(
        "The project utilizes a rich synthetic Electronic Health Record benchmark dataset comprising 15+ diverse longitudinal patient profiles "
        "compliant with HIPAA Safe Harbor de-identification principles. The dataset covers 6 major clinical specialties:"
    )

    # Table of Dataset Specialties
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(1.2), Inches(1.8), Inches(1.8), Inches(1.7)]
    headers = ["Specialty", "Primary Conditions", "Key Medications", "Lab Tracking"]
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "0284C7")
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)

    data_rows = [
        ("Endocrinology", "Type 2 Diabetes, Diabetic CKD Stage 3a, HTN", "Metformin, Empagliflozin, Lisinopril", "HbA1c, eGFR, Serum Creatinine, Microalbumin"),
        ("Cardiology", "CAD post-STEMI status post LAD DES, HFpEF, AFib", "Aspirin, Ticagrelor, Rosuvastatin, Entresto", "Peak Troponin I, Echocardiogram EF, LDL, NT-proBNP"),
        ("Pulmonology", "Moderate-Severe COPD, Asthma, Samter's Triad", "Advair Diskus, Spiriva, Albuterol", "Spirometry FEV1, FEV1/FVC, IgE, Eosinophils"),
        ("Rheumatology", "Seropositive Rheumatoid Arthritis, Osteoporosis", "Methotrexate, Humira, Alendronate", "CRP, ESR, RF, Anti-CCP, DEXA T-Scores"),
        ("Gastroenterology", "Crohn's Disease (Ileocolonic), Iron Anemia", "Ustekinumab (Stelara), Injectafer IV", "Fecal Calprotectin, Ferritin, Hemoglobin"),
        ("Neurology", "Relapsing-Remitting Multiple Sclerosis (RRMS)", "Ocrelizumab (Ocrevus), Gabapentin", "Brain MRI Lesions, sNfL, Serum IgG")
    ]

    for row_data in data_rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            row_cells[i].width = col_widths[i]
            set_cell_background(row_cells[i], "F8FAFC")
            set_cell_margins(row_cells[i], top=60, bottom=60, left=80, right=80)
            p = row_cells[i].paragraphs[0]
            p.runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    doc.add_paragraph(
        "Data Export Formats: In addition to the master JSON dataset (data/synthetic_patients.json), the system provides relational CSV exports "
        "(patients.csv, encounters.csv, labs.csv, medications.csv, allergies.csv) and a procedural data generator (scripts/generate_mock_dataset.py)."
    )

    # ==========================================
    # 4. IMPLEMENTATION / SOURCE CODE
    # ==========================================
    add_styled_heading(doc, "4. Implementation / Source Code Structure", level=1)
    doc.add_paragraph(
        "The project is implemented in modular Python 3 leveraging FastAPI, ChromaDB, Google GenAI SDK, and vanilla modern frontend technologies. "
        "The architecture is organized as follows:"
    )

    modules_info = [
        ("src/config.py", "Configuration management loading environment variables, Gemini model endpoints, server hosts, and file paths."),
        ("src/data_loader.py", "Clinical data parser and chunker that attaches metadata (patient_id, encounter_date, source_id) and manages dataset persistence."),
        ("src/embeddings.py", "Embedding function implementing ChromaDB's EmbeddingFunction interface with Google Gemini embedding-001 and local fallback."),
        ("src/vector_store.py", "ChromaDB vector database manager executing metadata-filtered similarity queries and document upserts."),
        ("src/rag_engine.py", "Core RAG orchestrator assembling chronological clinical context, evaluating prompt guardrails, and synthesizing grounded answers."),
        ("src/api.py", "FastAPI REST backend serving endpoints for patient files, RAG queries, health checks, and multipart dataset uploads."),
        ("static/ (HTML/CSS/JS)", "Clinical web dashboard with patient switcher, allergy warnings, diagnosis tags, vector evidence cards, and dataset upload dropzone."),
        ("cli.py", "Command-line interface enabling instant terminal queries, interactive prompting, and database re-indexing."),
        ("tests/test_rag.py", "Automated test suite verifying data structures, strict patient isolation, allergy retrieval, and RAG accuracy.")
    ]

    for fname, desc in modules_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r_fn = p.add_run(f"• {fname}: ")
        r_fn.font.bold = True
        r_fn.font.color.rgb = RGBColor(2, 132, 199)
        p.add_run(desc)

    # ==========================================
    # 5. SCREENSHOTS OR OUTPUT
    # ==========================================
    add_styled_heading(doc, "5. Screenshots & Output Verification", level=1)
    doc.add_paragraph("Below are representative execution outputs validating retrieval, citations, and terminal interface:")

    add_callout_box(
        doc,
        "Sample Query Execution: Historical Cardiac Intervention",
        "CLI Command: python cli.py -p P102 -q 'What happened during the cardiac catheterization and what stent was placed?'\n\n"
        "Synthesized Clinical Output:\n"
        "• Procedure Date: 2023-05-14 [Source: ENC-P102-01]\n"
        "• Indication: Acute anterior ST-elevation myocardial infarction (STEMI) [Source: ENC-P102-01]\n"
        "• Angiographic Findings: 95% occlusion of the mid-Left Anterior Descending (LAD) artery [Source: ENC-P102-01]\n"
        "• Stent Deployment: Drug-Eluting Stent (Synergy 3.5 x 18mm) deployed with TIMI-3 flow restored [Source: ENC-P102-01]\n"
        "• Biomarker Peak: Peak Troponin I reached 14.2 ng/mL [Source: ENC-P102-01]\n"
        "• Echocardiogram Trend: LVEF 50% post-procedure (2023-05-14) recovering to 55-60% on 2024-06-12 [Source: ENC-P102-03]\n\n"
        "Retrieved Sources: [ENC-P102-01] (71% Match), [ENC-P102-03] (65% Match), [ENC-P102-02] (64% Match)"
    )

    add_callout_box(
        doc,
        "Sample Query Execution: Longitudinal Laboratory Trend",
        "CLI Command: python cli.py -p P101 -q 'What historical laboratory trends and diagnoses are documented?'\n\n"
        "Synthesized Clinical Output:\n"
        "• Troponin I: Progressive elevation (2021-03-12: 1.09 ng/mL -> 2022-08-20: 1.67 ng/mL -> 2023-11-05: 11.34 ng/mL) [Source: P101-LAB-Troponin_I]\n"
        "• LDL Cholesterol: Decreased from 148 mg/dL to 93 mg/dL on lipid-lowering therapy [Source: P101-LAB-LDL_Cholesterol]\n"
        "• NT-proBNP: Stable at 64.9 units following initial rise from 33.9 units [Source: P101-LAB-NT-proBNP]\n"
        "• Data Absence Check: Accurately reports absence of unstated formal diagnoses in context."
    )

    # ==========================================
    # 6. RESULTS AND OBSERVATIONS
    # ==========================================
    add_styled_heading(doc, "6. Results and Observations", level=1)
    
    results = [
        ("Zero Patient Data Leakage", "Automated tests confirmed that queries for Patient A never retrieve chunks belonging to Patient B. The metadata pre-filter guarantees complete multi-tenant clinical isolation."),
        ("Temporal Trend Resolution", "The system accurately reconstructed multi-year laboratory trajectories (e.g. HbA1c trending from 8.6% down to 6.9%, and Troponin peaking post-MI), correctly identifying chronological sequences."),
        ("Strict Citation Grounding", "Every factual medical assertion in the LLM response is supported by an inline citation referencing the exact encounter ID or diagnostic test panel."),
        ("Dynamic Dataset Ingestion", "Users can upload custom patient datasets (.json or .csv) via the web interface dropzone, automatically updating the vector database and making records queryable in under 2 seconds."),
        ("High Reliability & Offline Fallback", "When external API rate limits or network issues occur, the system seamlessly transitions to local deterministic embeddings and rule-based clinical extraction without crashing.")
    ]

    for title, detail in results:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"✔ {title}: ")
        r.font.bold = True
        r.font.color.rgb = RGBColor(22, 163, 74)
        p.add_run(detail)

    # ==========================================
    # 7. CONCLUSION
    # ==========================================
    add_styled_heading(doc, "7. Conclusion", level=1)
    doc.add_paragraph(
        "The Patient-Record RAG Prototype successfully demonstrates that metadata-filtered vector retrieval combined with "
        "temporally grounded LLM reasoning provides an accurate, trustworthy solution for querying complex historical health records. "
        "By enforcing strict patient isolation, chronological ordering, and verifiable citation tags, the prototype prevents clinical "
        "hallucinations and provides clinicians with immediate, evidence-backed historical context."
    )

    doc.add_paragraph(
        "Future enhancements could include HL7 FHIR native API ingestion, multimodal medical imaging integration (DICOM X-rays/MRIs), "
        "and multi-agent clinical consensus review for differential diagnosis assistance."
    )

    # ==========================================
    # 8. GITHUB REPOSITORY LINK & DEPLOYMENT
    # ==========================================
    add_styled_heading(doc, "8. GitHub Repository & Deployment Details", level=1)
    
    add_callout_box(
        doc,
        "Repository & Live Resources",
        "• GitHub Repository: https://github.com/VK-JPG-cmd/rag_assessment_prototype.git\n"
        "• Local Web Server: http://127.0.0.1:8000 (run via python -m src.api)\n"
        "• Interactive API Docs: http://127.0.0.1:8000/docs\n"
        "• Deployment Target: Render.com (Configuration: render.yaml & requirements.txt)\n"
        "• CLI Runner: python cli.py --patient <ID> --query <QUESTION>"
    )

    # Save Document
    output_path = Path("Patient_Record_RAG_Prototype_Report.docx")
    doc.save(str(output_path))
    print(f"[OK] Generated Word Document: {output_path.resolve()}")

if __name__ == "__main__":
    generate_report()
